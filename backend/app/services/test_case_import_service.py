"""Import authored test cases from CSV or Excel (W13 TCL-04).

A dry run first (nothing added, nothing flushed), then the apply re-reads the
same file. Every case lands as a Draft through the ONE create path
(``test_management_service.create_managed_test_case``: membership, suite
resolution, version 1, the audit row), one SAVEPOINT per row. Imported text is
untrusted: caps are checked BEFORE parsing (bytes while reading, an archive's
expanded size before ``load_workbook``, rows while iterating), cells are kept
as typed (a leading ``=`` is data; every export escapes it again) and only the
one ``'`` an export adds before a formula character is removed. Log lines carry
counts and ids only, never a cell, a title or the file's content.

Stage-only: the router commits (once, when at least one case was created).
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any, Generator, Optional

import structlog
from fastapi import HTTPException, UploadFile

from app.core.test_library_vocab import (
    FORMULA_PREFIXES,
    IMPORT_HEADERS,
    IMPORT_IGNORED_HEADERS,
    IMPORT_MAX_BYTES,
    IMPORT_MAX_ROWS,
    IMPORT_MAX_XLSX_UNCOMPRESSED,
    IMPORT_NO_TITLE_COLUMN,
    IMPORT_NOT_A_FILE,
    IMPORT_TOO_LARGE,
    IMPORT_TOO_MANY_ROWS,
    IMPORT_UNREADABLE,
    IMPORT_XLSX_TOO_BIG,
)
from app.services.test_management_audit_service import activity_summarised_elsewhere

logger = structlog.get_logger(__name__)

_READ_CHUNK = 64 * 1024
_XML_INVALID = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")      # = test_management_exports._strip_control_chars
_FILE_NAME_UNSAFE = re.compile(r"[\x00-\x1f\x7f]")
_XLSX_MAGIC = b"PK\x03\x04"


@dataclass
class ParsedImport:
    """A parsed upload: the format, the header report, and the data rows by column key."""

    format: str
    recognised: list[str] = field(default_factory=list)
    ignored: list[str] = field(default_factory=list)
    unknown: list[str] = field(default_factory=list)
    #: (spreadsheet row number, {field or "status": raw cell text})
    rows: list[tuple[int, dict[str, str]]] = field(default_factory=list)


async def read_bounded(upload: UploadFile, max_bytes: int = IMPORT_MAX_BYTES) -> bytes:
    """The upload's bytes; 413 ``IMPORT_TOO_LARGE`` as soon as it is over the cap (never read past it)."""
    chunks: list[bytes] = []
    total = 0
    while True:
        chunk = await upload.read(_READ_CHUNK)
        if not chunk:
            break
        total += len(chunk)
        if total > max_bytes:
            raise HTTPException(status_code=413, detail=IMPORT_TOO_LARGE)
        chunks.append(chunk)
    return b"".join(chunks)


def safe_file_name(raw: Optional[str]) -> str:
    """The upload's base name, control characters dropped, at most 200 characters."""
    name = re.split(r"[\\/]", raw or "")[-1]
    name = _FILE_NAME_UNSAFE.sub("", name).strip()
    return (name or "upload")[:200]


def _cell_text(value: Any) -> str:
    """One cell as text: a number as its digits, a date as an ISO date, a boolean as Yes / No."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    return str(value)


def _csv_rows(content: bytes) -> Generator[list[str], None, None]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise HTTPException(status_code=422, detail=IMPORT_UNREADABLE) from None
    try:
        delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t").delimiter
    except csv.Error:
        delimiter = ","
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    try:
        for row in reader:
            yield list(row)
    except csv.Error:                       # a cell over csv.field_size_limit (untouched), a NUL ...
        raise HTTPException(status_code=422, detail=IMPORT_UNREADABLE) from None


def _xlsx_rows(content: bytes) -> Generator[list[str], None, None]:
    if not content.startswith(_XLSX_MAGIC):
        raise HTTPException(status_code=400, detail=IMPORT_NOT_A_FILE)
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            expanded = sum(member.file_size for member in archive.infolist())
    except (zipfile.BadZipFile, ValueError, OSError):
        raise HTTPException(status_code=422, detail=IMPORT_UNREADABLE) from None
    if expanded > IMPORT_MAX_XLSX_UNCOMPRESSED:
        raise HTTPException(status_code=422, detail=IMPORT_XLSX_TOO_BIG)
    from openpyxl import load_workbook      # defusedxml is installed: openpyxl parses through it

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
    except Exception:  # noqa: BLE001 -- any malformed workbook is the one sentence
        raise HTTPException(status_code=422, detail=IMPORT_UNREADABLE) from None
    try:
        sheet = workbook.worksheets[0] if workbook.worksheets else None
        if sheet is None:
            return
        for values in sheet.iter_rows(values_only=True):
            yield [_cell_text(v) for v in values]
    except HTTPException:
        raise
    except Exception:  # noqa: BLE001
        raise HTTPException(status_code=422, detail=IMPORT_UNREADABLE) from None
    finally:
        workbook.close()


def _header_key(text: str) -> str:
    return " ".join(text.split()).lower()


def unquote_export_cell(text: str) -> tuple[str, bool]:
    """The exact inverse of ``_excel_text``: ONE leading ``'`` before a formula character is removed."""
    if text.startswith("'") and text[1:].lstrip().startswith(FORMULA_PREFIXES):
        return text[1:], True
    return text, False


def _format_of(file_name: str) -> str:
    lowered = (file_name or "").lower()
    if lowered.endswith(".csv"):
        return "csv"
    if lowered.endswith(".xlsx"):
        return "xlsx"
    raise HTTPException(status_code=400, detail=IMPORT_NOT_A_FILE)


def _columns(header: list[str], parsed: ParsedImport) -> dict[int, str]:
    """Column index -> field (or ``status``); fills the header report. The first of a repeated header wins."""
    columns: dict[int, str] = {}
    taken: set[str] = set()
    for index, raw in enumerate(header):
        written = raw.strip()
        if not written:
            continue
        key = _header_key(written)
        target = IMPORT_HEADERS.get(key) or ("status" if key == "status" else None)
        if target is not None and target not in taken:
            columns[index] = target
            taken.add(target)
            (parsed.ignored if target == "status" else parsed.recognised).append(written)
        elif key in IMPORT_IGNORED_HEADERS and target is None:
            parsed.ignored.append(written)
        else:
            parsed.unknown.append(written)
    if "title" not in taken:
        raise HTTPException(status_code=422, detail=IMPORT_NO_TITLE_COLUMN)
    return columns


def parse_upload(file_name: str, content: bytes) -> ParsedImport:
    """Pure: the format, the header report and every non-empty data row (at most ``IMPORT_MAX_ROWS``).

    400 not a .csv / .xlsx (or an .xlsx that is no zip); 422 unreadable, a
    workbook expanding past 50 MB, no Title column, more than 2,000 rows.
    """
    fmt = _format_of(file_name)
    parsed = ParsedImport(format=fmt)
    rows = _csv_rows(content) if fmt == "csv" else _xlsx_rows(content)
    columns: Optional[dict[int, str]] = None
    try:
        for number, cells in enumerate(rows, start=1):
            if not any((c or "").strip() for c in cells):
                continue                                   # a fully empty row is skipped, not counted
            if columns is None:
                columns = _columns(cells, parsed)
                continue
            if len(parsed.rows) >= IMPORT_MAX_ROWS:
                raise HTTPException(status_code=422, detail=IMPORT_TOO_MANY_ROWS)
            values = {target: _XML_INVALID.sub("", cells[index]) for index, target in columns.items()
                      if index < len(cells) and cells[index]}
            parsed.rows.append((number, values))
    finally:
        rows.close()                                       # the workbook closes with its generator
    if columns is None:
        raise HTTPException(status_code=422, detail=IMPORT_NO_TITLE_COLUMN)
    return parsed


# ── one row -> the create payload's fields, or why not ───────────────────────

_COLUMN_LABELS = {
    "title": "Title", "test_type": "Type", "priority": "Priority", "severity": "Severity",
    "feature_area": "Feature Area", "objective": "Objective", "preconditions": "Preconditions",
    "steps": "Steps", "expected_result": "Expected Result", "is_automated": "Automated",
    "automation_status": "Automation Status", "tags": "Tags", "suite_name": "Suite",
    "description": "Description", "test_data": "Test Data", "status": "Status",
}
_TYPE_RE = re.compile(r"^[a-z0-9 _-]{1,50}$")
_YES = {"yes", "true", "1", "y"}
_NO = {"no", "false", "0", "n", ""}
_STEP_SPLIT = re.compile(r";\s*(?=step\s+\d+\s*:)", re.IGNORECASE)
_STEP_HEAD = re.compile(r"^\s*step\s+\d+\s*:\s*", re.IGNORECASE)
_SEVERITY_MAX = 20                                         # managed_test_cases.severity is VARCHAR(20)


@dataclass
class RowCheck:
    row: int
    fields: dict[str, Any]
    reasons: list[str]
    warnings: list[str]
    status: Optional[str] = None                           # a non-draft Status the file named (not applied)


def parse_steps(text: str) -> list[dict]:
    """``Step N: …`` segments (the export's own format), else one step per non-empty line."""
    text = text.strip()
    if not text:
        return []
    parts = _STEP_SPLIT.split(text) if _STEP_HEAD.match(text) else text.splitlines()
    actions = [_STEP_HEAD.sub("", part, count=1).strip() for part in parts]
    return [{"step_number": n, "action": a, "expected_result": ""} for n, a in enumerate((a for a in actions if a), 1)]


def check_row(row: int, raw: dict[str, str]) -> RowCheck:
    """Validate one data row (vocabularies, lengths, the export's ``'``); never raises."""
    from app.core import test_library_vocab as v

    check = RowCheck(row=row, fields={}, reasons=[], warnings=[])
    cells: dict[str, str] = {}
    for key, text in raw.items():
        text, unquoted = unquote_export_cell(text)
        if unquoted:
            check.warnings.append(v.IMPORT_WARN_QUOTE.format(column=_COLUMN_LABELS[key]))
        if len(text) > v.IMPORT_MAX_CELL_CHARS:
            check.reasons.append(v.IMPORT_ROW_CELL_TOO_LONG.format(column=_COLUMN_LABELS[key]))
            continue
        cells[key] = text
    title = cells.get("title", "").strip()
    if not 3 <= len(title) <= 500:
        check.reasons.append(v.IMPORT_ROW_TITLE)
    check.fields["title"] = title
    priority = cells.get("priority", "").strip().lower() or "medium"
    priority = v.PRIORITY_ALIASES.get(priority, priority)
    if priority not in v.CASE_PRIORITIES:
        check.reasons.append(v.IMPORT_ROW_PRIORITY)
    check.fields["priority"] = priority
    severity = cells.get("severity", "").strip().lower() or "major"
    if severity not in v.CASE_SEVERITIES:
        if len(severity) > _SEVERITY_MAX:
            check.reasons.append(v.IMPORT_SAVE_FAILED.format(detail="Severity is longer than 20 characters."))
        else:
            check.warnings.append(v.IMPORT_WARN_SEVERITY.format(value=cells.get("severity", "").strip()))
    check.fields["severity"] = severity
    test_type = cells.get("test_type", "").strip().lower() or "functional"
    if not _TYPE_RE.match(test_type):
        check.reasons.append(v.IMPORT_ROW_TYPE)
    check.fields["test_type"] = test_type
    automated = cells.get("is_automated", "").strip().lower()
    if automated not in _YES | _NO:
        check.reasons.append(v.IMPORT_ROW_AUTOMATED)
    check.fields["is_automated"] = automated in _YES
    automation_status = cells.get("automation_status", "").strip().lower() or "not_automated"
    if automation_status not in v.AUTOMATION_STATUSES:
        check.reasons.append(v.IMPORT_ROW_AUTOMATION_STATUS)
    check.fields["automation_status"] = automation_status
    tags = list(dict.fromkeys(t.strip() for t in re.split(r"[,;]", cells.get("tags", "")) if t.strip()))
    if len(tags) > 20 or any(len(t) > 100 for t in tags):
        check.reasons.append(v.IMPORT_ROW_TAGS)
    check.fields["tags"] = tags or None
    steps = parse_steps(cells.get("steps", ""))
    check.fields["steps"] = steps or None
    for key in ("feature_area", "objective", "preconditions", "expected_result", "suite_name", "description",
                "test_data"):
        text = cells.get(key, "")
        check.fields[key] = text if text.strip() else None
    status = cells.get("status", "").strip()
    if status and status.lower() != "draft":
        check.status = status
        check.warnings.append(v.IMPORT_WARN_STATUS.format(value=status))
    return check


# ── the import ────────────────────────────────────────────────────────────────


def _fingerprint(title: Any, objective: Any, steps: Any, expected_result: Any) -> str:
    from types import SimpleNamespace

    from app.services.duplicate_detection_service import compute_dup_fingerprint

    return compute_dup_fingerprint(SimpleNamespace(  # type: ignore[arg-type]
        title=title, objective=objective, steps=steps, expected_result=expected_result,
    ))


async def _existing_fingerprints(db: Any, project_id: Any) -> dict[str, tuple[Any, str]]:
    """ONE read: every duplicate-eligible case of the project, fingerprint -> (id, title)."""
    from sqlalchemy import select

    from app.models.postgres import ManagedTestCase as M
    from app.services.duplicate_detection_service import DUPLICATE_ELIGIBLE_STATES

    stmt = select(M.id, M.title, M.objective, M.steps, M.expected_result, M.dup_fingerprint).where(
        M.project_id == project_id, M.status.in_(DUPLICATE_ELIGIBLE_STATES),
    )
    known: dict[str, tuple[Any, str]] = {}
    for r in (await db.execute(stmt)).all():
        fp = r.dup_fingerprint or _fingerprint(r.title, r.objective, r.steps, r.expected_result)
        known.setdefault(fp, (r.id, r.title))
    return known


def _validation_detail(exc: Exception) -> str:
    from pydantic import ValidationError

    if isinstance(exc, HTTPException):
        return str(exc.detail)
    if isinstance(exc, ValidationError) and exc.errors():
        first = exc.errors()[0]
        where = _COLUMN_LABELS.get(str(first.get("loc", ("",))[0]), "A value")
        return f"{where}: {first.get('msg', 'not valid')}."
    return "the database refused one of its values."


def _row_result(check: RowCheck, outcome: str, **extra: Any) -> dict:
    return {"row": check.row, "outcome": outcome, "title": check.fields.get("title") or None,
            "reasons": list(check.reasons), "warnings": list(check.warnings),
            "duplicate_of": extra.get("duplicate_of"), "case_id": extra.get("case_id")}


async def import_cases(db: Any, *, project_id: Any, actor: Any, file_name: str, content: bytes,
                       dry_run: bool) -> dict:
    """A ``TestCaseImportResponse`` body. A dry run writes nothing (no add, no flush); an apply
    creates each valid row in its own SAVEPOINT and stages ONE ``test_case.imported`` row."""
    from app.core.test_library_vocab import IMPORT_DUPLICATE_EXISTING, IMPORT_DUPLICATE_ROW, IMPORT_SAVE_FAILED
    from app.models.schemas import ManagedTestCaseCreate

    name = safe_file_name(file_name)
    parsed = parse_upload(name, content)
    known = await _existing_fingerprints(db, project_id)
    seen_rows: dict[str, int] = {}
    results: list[dict] = []
    to_create: list[tuple[RowCheck, ManagedTestCaseCreate, str]] = []
    for row, raw in parsed.rows:
        check = check_row(row, raw)
        if check.reasons:
            results.append(_row_result(check, "invalid"))
            continue
        try:
            payload = ManagedTestCaseCreate(project_id=project_id, **check.fields)
        except Exception as exc:  # noqa: BLE001 -- a schema limit is this row's reason
            check.reasons.append(IMPORT_SAVE_FAILED.format(detail=_validation_detail(exc)))
            results.append(_row_result(check, "invalid"))
            continue
        fp = _fingerprint(payload.title, payload.objective, check.fields["steps"], payload.expected_result)
        if fp in known:
            case_id, title = known[fp]
            check.reasons.append(IMPORT_DUPLICATE_EXISTING)
            results.append(_row_result(check, "duplicate", duplicate_of={"case_id": case_id, "title": title}))
        elif fp in seen_rows:
            check.reasons.append(IMPORT_DUPLICATE_ROW.format(row=seen_rows[fp]))
            results.append(_row_result(check, "duplicate", duplicate_of={"row": seen_rows[fp]}))
        else:
            seen_rows[fp] = row
            results.append(_row_result(check, "create"))
            to_create.append((check, payload, fp))
    if not dry_run:
        await _apply(db, project_id=project_id, actor=actor, file_name=name, rows=to_create, results=results)
    return _response(project_id, dry_run, name, parsed, results, attempted=len(to_create))


@activity_summarised_elsewhere
async def _apply(db: Any, *, project_id: Any, actor: Any, file_name: str,
                 rows: list[tuple[RowCheck, Any, str]], results: list[dict]) -> None:
    """Create each row in its own SAVEPOINT (a failure leaves the others), then ONE activity row."""
    from app.core.test_library_vocab import IMPORT_SAVE_FAILED
    from app.services.activity.service import ActorRef, record as record_activity
    from app.services.test_management_audit_service import audit_event
    from app.services.test_management_service import create_managed_test_case

    by_row = {r["row"]: r for r in results}
    created = 0
    for check, payload, fp in rows:
        result = by_row[check.row]
        try:
            async with db.begin_nested():
                case = await create_managed_test_case(db, payload, actor)
                case.dup_fingerprint = fp
                if check.status is not None:
                    await audit_event(
                        db, "test_case", case.id, project_id, "imported", actor,
                        details=f"Imported from {file_name} (row {check.row}): its Status "
                                f"'{check.status}' was not applied",
                    )
                await db.flush()
        except Exception as exc:  # noqa: BLE001 -- the row is reported, the others go on
            logger.warning("test_case_import_row_failed", row=check.row, error_type=type(exc).__name__)
            result.update(outcome="invalid",
                          reasons=[IMPORT_SAVE_FAILED.format(detail=_validation_detail(exc))])
            continue
        created += 1
        result.update(outcome="created", case_id=case.id)
    if created:
        invalid = sum(1 for r in results if r["outcome"] == "invalid")
        duplicate = sum(1 for r in results if r["outcome"] == "duplicate")
        await record_activity(
            db,
            project_id=project_id,
            event_type="test_case.imported",
            actor=ActorRef.from_user(actor),
            entity_id=project_id,
            entity_type="test_library",
            entity_label=file_name[:200],
            context={"count": created, "duplicate": duplicate, "invalid": invalid},
        )


def _response(project_id: Any, dry_run: bool, file_name: str, parsed: ParsedImport, results: list[dict],
              *, attempted: int) -> dict:
    counts = {outcome: sum(1 for r in results if r["outcome"] == outcome)
              for outcome in ("create", "duplicate", "invalid", "created")}
    return {
        "project_id": project_id,
        "dry_run": dry_run,
        "file_name": file_name,
        "format": parsed.format,
        "headers": {"recognised": parsed.recognised, "ignored": parsed.ignored, "unknown": parsed.unknown},
        "summary": {"rows": len(results), "create": attempted if not dry_run else counts["create"],
                    "duplicate": counts["duplicate"], "invalid": counts["invalid"], "created": counts["created"]},
        "rows": results,
    }
