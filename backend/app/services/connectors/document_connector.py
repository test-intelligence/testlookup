"""Document upload connector — reads uploaded files from MinIO and extracts text."""
from __future__ import annotations

import io
import time
from typing import Optional

import structlog

from app.core.config import settings
from app.services.connectors.base import (
    ConnectorFetchError,
    FetchedContent,
    KnowledgeConnectorBase,
)

logger = structlog.get_logger(__name__)


class DocumentConnector(KnowledgeConnectorBase):
    """Handles uploaded_document source type — downloads from MinIO and extracts text."""

    connector_type = "uploaded_document"

    async def test_connection(self) -> dict:
        t_start = time.monotonic()
        try:
            from app.db.storage import get_storage_provider
            storage = get_storage_provider()
            # Try listing the knowledge docs bucket to verify connectivity
            await storage.list_objects(
                prefix="", bucket=settings.KNOWLEDGE_DOCS_BUCKET
            )
            latency = int((time.monotonic() - t_start) * 1000)
            return {
                "success": True,
                "latency_ms": latency,
                "error": None,
                "detail": f"Storage accessible (bucket: {settings.KNOWLEDGE_DOCS_BUCKET})",
            }
        except Exception as exc:
            latency = int((time.monotonic() - t_start) * 1000)
            return {
                "success": False,
                "latency_ms": latency,
                "error": str(exc)[:200],
                "detail": "Cannot access storage backend",
            }

    async def fetch_content(
        self,
        canonical_url: str,
        external_id: Optional[str] = None,
    ) -> FetchedContent:
        """Download file from MinIO and extract text content.

        canonical_url is the MinIO object key for uploaded documents.
        """
        storage_key = canonical_url
        filename = storage_key.rsplit("/", 1)[-1] if "/" in storage_key else storage_key

        try:
            from app.db.storage import get_storage_provider
            storage = get_storage_provider()
            file_bytes = await storage.get_object_content(
                key=storage_key, bucket=settings.KNOWLEDGE_DOCS_BUCKET
            )
        except FileNotFoundError:
            raise ConnectorFetchError(f"Document not found in storage: {storage_key}")
        except Exception as exc:
            raise ConnectorFetchError(
                f"Failed to download document from storage: {exc}", retryable=True
            )

        ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
        try:
            raw_text = _extract_text(file_bytes, ext, filename)
        except ConnectorFetchError:
            raise
        except Exception as exc:
            raise ConnectorFetchError(f"Text extraction failed for {filename}: {exc}")

        if not raw_text.strip():
            raise ConnectorFetchError(f"No text content extracted from {filename}")

        return FetchedContent(
            raw_text=raw_text,
            content_hash=FetchedContent.compute_hash(raw_text),
            title=f"Document: {filename}",
            source_url=storage_key,
            metadata={
                "connector": "document",
                "filename": filename,
                "file_type": ext,
                "size_bytes": len(file_bytes),
            },
        )


def _extract_text(file_bytes: bytes, ext: str, filename: str) -> str:
    """Route to appropriate text extractor based on file extension."""
    if ext == "pdf":
        return _extract_pdf(file_bytes)
    if ext in ("docx", "doc"):
        return _extract_docx(file_bytes)
    if ext in ("md", "markdown", "rst", "txt"):
        return _extract_plaintext(file_bytes, filename)
    raise ConnectorFetchError(f"Unsupported file type: .{ext}")


def _extract_pdf(file_bytes: bytes) -> str:
    """Extract text from PDF using pypdf."""
    try:
        from pypdf import PdfReader
    except ImportError:
        raise ConnectorFetchError(
            "pypdf is not installed — required for PDF extraction (pip install pypdf)"
        )

    reader = PdfReader(io.BytesIO(file_bytes))
    parts: list[str] = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text()
        if text and text.strip():
            parts.append(f"--- Page {i + 1} ---\n{text.strip()}")
    return "\n\n".join(parts)


def _extract_docx(file_bytes: bytes) -> str:
    """Extract text from DOCX using python-docx."""
    try:
        from docx import Document
    except ImportError:
        raise ConnectorFetchError(
            "python-docx is not installed — required for DOCX extraction (pip install python-docx)"
        )

    doc = Document(io.BytesIO(file_bytes))
    parts: list[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        style_name = (para.style.name or "").lower() if para.style else ""
        if "heading 1" in style_name:
            parts.append(f"# {text}")
        elif "heading 2" in style_name:
            parts.append(f"## {text}")
        elif "heading 3" in style_name:
            parts.append(f"### {text}")
        elif "list" in style_name:
            parts.append(f"- {text}")
        else:
            parts.append(text)

    # Extract tables
    for table in doc.tables:
        table_rows: list[str] = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            table_rows.append("| " + " | ".join(cells) + " |")
        if table_rows:
            parts.append("\n".join(table_rows))

    return "\n\n".join(parts)


def _extract_plaintext(file_bytes: bytes, filename: str) -> str:
    """Extract text from plain text, Markdown, or RST files."""
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            return file_bytes.decode(encoding)
        except (UnicodeDecodeError, ValueError):
            continue
    raise ConnectorFetchError(f"Cannot decode text file {filename} — unsupported encoding")
