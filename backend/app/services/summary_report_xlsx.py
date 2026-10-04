"""VIZ-607 — the Summary Report as an Excel workbook, charts included.

One sheet per part of the report, each with its data and, where it reads as
a chart, a native Excel chart over that data (``openpyxl.chart``), so the
numbers in a circulated workbook stay inspectable and re-chartable:

* **Report context** -- the context block every export carries (VIZ-308:
  Project, Release, Test Suite, Window, Aggregation, basis, Generated, and
  the "N of M" line), from the response's ``meta``.
* **Totals** -- the headline counts and rates, with a status pie.
* **Suites** -- the per-suite breakdown, with a pass-rate bar chart.
* **Top failing tests** -- with a failures bar chart.

Every text cell goes through ``_excel_text``: stored names are untrusted, and
a cell that starts with ``=``/``+``/``-``/``@`` would otherwise be a formula
(the rule ``test_management_exports._excel_text`` applies to its workbook).
"""
from __future__ import annotations

import io
import re
from typing import Any

from app.services.privacy_service import sanitize_for_report
from app.services.report_context import context_fields, has_context, n_of_m_text

_XML_INVALID = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")

#: Bars per chart; the sheet itself holds every row.
CHART_ROWS = 15


def _excel_text(value: Any) -> str:
    text = _XML_INVALID.sub("", sanitize_for_report(str(value if value is not None else "")))
    return "'" + text if text.lstrip().startswith(_FORMULA_START) else text


def _num(value: Any) -> Any:
    return value if isinstance(value, (int, float)) and not isinstance(value, bool) else None


def render_summary_report_xlsx(payload: dict) -> bytes:
    """The workbook's bytes. Synchronous and CPU-bound: call it off the event loop."""
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, PieChart, Reference
    from openpyxl.styles import Font

    bold = Font(bold=True)
    wb = Workbook()

    # ── Report context ──────────────────────────────────────────────────
    ctx = wb.active
    ctx.title = "Report context"
    ctx.append(["TestLookup — Summary Report"])
    ctx["A1"].font = Font(bold=True, size=14)
    meta = payload.get("meta")
    mode = "All runs in window" if payload.get("mode", "window") == "window" else "Latest run per suite"
    if has_context(meta):
        for field in context_fields(meta, aggregation=mode):
            ctx.append([_excel_text(field.label), _excel_text(field.value)])
        ctx.append([])
        ctx.append([_excel_text(n_of_m_text(meta))])
    else:
        ctx.append(["Project", _excel_text(payload.get("project_name") or "—")])
        ctx.append(["Window (days)", _num(payload.get("window_days"))])
        ctx.append(["Aggregation", mode])
        ctx.append(["Generated", _excel_text(payload.get("generated_at") or "")])
    for row in ctx.iter_rows(min_row=2, max_col=1):
        row[0].font = bold
    ctx.column_dimensions["A"].width = 22
    ctx.column_dimensions["B"].width = 90

    # ── Totals, with a status pie ───────────────────────────────────────
    totals = payload.get("totals") or {}
    tot = wb.create_sheet("Totals")
    tot.append(["Status", "Tests"])
    statuses = [("Passed", "passed"), ("Failed", "failed"), ("Broken", "broken"), ("Skipped", "skipped")]
    for label, key in statuses:
        tot.append([label, int(totals.get(key) or 0)])
    tot.append([])
    tot.append(["Measure", "Value"])
    for label, key in [
        ("Total tests", "total_test_cases"),
        ("Evaluated", "evaluated"),
        ("Pass %", "pass_rate_pct"),
        ("Fail %", "fail_rate_pct"),
        ("Skip %", "skip_rate_pct"),
        ("Broken %", "broken_rate_pct"),
        ("Weighted pass %", "weighted_pass_rate_pct"),
    ]:
        tot.append([label, _num(totals.get(key))])
    tot.append(["Pass-rate basis", _excel_text(totals.get("pass_rate_basis_label") or "")])
    tot.append(["Runs", _num(payload.get("run_count"))])
    tot.append(["Flaky tests", _num(payload.get("flaky_test_count"))])
    for cell in (tot["A1"], tot["B1"], tot["A7"], tot["B7"]):
        cell.font = bold
    tot.column_dimensions["A"].width = 20
    if sum(int(totals.get(key) or 0) for _, key in statuses) > 0:
        pie = PieChart()
        pie.title = "Status breakdown"
        pie.add_data(Reference(tot, min_col=2, min_row=1, max_row=5), titles_from_data=True)
        pie.set_categories(Reference(tot, min_col=1, min_row=2, max_row=5))
        tot.add_chart(pie, "D2")

    # ── Suites, with a pass-rate bar chart ──────────────────────────────
    suites = payload.get("suites") or []
    su = wb.create_sheet("Suites")
    su.append(["Suite", "Total", "Passed", "Failed", "Skipped", "Broken", "Pass %"])
    for cell in su[1]:
        cell.font = bold
    for s in suites:
        su.append([
            _excel_text(s.get("suite_name")),
            _num(s.get("total")), _num(s.get("passed")), _num(s.get("failed")),
            _num(s.get("skipped")), _num(s.get("broken")), _num(s.get("pass_rate_pct")),
        ])
    su.column_dimensions["A"].width = 40
    if suites:
        bars = BarChart()
        bars.type = "bar"
        bars.title = "Pass rate by suite"
        bars.y_axis.title = "Pass %"
        bars.y_axis.scaling.min = 0
        bars.y_axis.scaling.max = 100
        last = 1 + min(len(suites), CHART_ROWS)
        bars.add_data(Reference(su, min_col=7, min_row=1, max_row=last), titles_from_data=True)
        bars.set_categories(Reference(su, min_col=1, min_row=2, max_row=last))
        bars.height = max(7.5, 0.6 * (last - 1))
        su.add_chart(bars, "I2")

    # ── Top failing tests, with a failures bar chart ────────────────────
    top = payload.get("top_failing_tests") or []
    tf = wb.create_sheet("Top failing tests")
    tf.append(["Test", "Suite", "Class", "Failures"])
    for cell in tf[1]:
        cell.font = bold
    for t in top:
        tf.append([
            _excel_text(t.get("test_name")), _excel_text(t.get("suite_name") or ""),
            _excel_text(t.get("class_name") or ""), _num(t.get("failures")),
        ])
    tf.column_dimensions["A"].width = 60
    tf.column_dimensions["B"].width = 24
    tf.column_dimensions["C"].width = 30
    if top:
        bars = BarChart()
        bars.type = "bar"
        bars.title = "Top failing tests"
        bars.y_axis.title = "Failures"
        bars.y_axis.scaling.min = 0
        last = 1 + min(len(top), CHART_ROWS)
        bars.add_data(Reference(tf, min_col=4, min_row=1, max_row=last), titles_from_data=True)
        bars.set_categories(Reference(tf, min_col=1, min_row=2, max_row=last))
        bars.height = max(7.5, 0.6 * (last - 1))
        tf.add_chart(bars, "F2")

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
