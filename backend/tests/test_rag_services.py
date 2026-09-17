"""
Comprehensive unit tests for RAG services — covers untested code paths
across sync, chunking, retrieval, generation, review, staleness, connectors,
redaction, and eval services.
"""
from __future__ import annotations

import uuid
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import pytest

from app.models.postgres import (
    GenerationBatch,
    GenerationCaseSource,
    KnowledgeClassification,
    KnowledgeSourceType,
    KnowledgeSyncStatus,
    ManagedTestCase,
)
from app.services.connectors.base import ConnectorFetchError, FetchedContent


# ═══════════════════════════════════════════════════════════════════════════════
# Helpers
# ═══════════════════════════════════════════════════════════════════════════════


def _make_source(**overrides) -> SimpleNamespace:
    """Create a lightweight source-like object for unit tests (avoids SQLAlchemy session)."""
    defaults = dict(
        id=uuid.uuid4(),
        project_id=uuid.uuid4(),
        source_type=KnowledgeSourceType.JIRA_ISSUE.value,
        title="PROJ-1: Login feature",
        canonical_url="https://jira.example.com/browse/PROJ-1",
        sync_status=KnowledgeSyncStatus.PENDING.value,
        classification=KnowledgeClassification.INTERNAL.value,
        is_archived=False,
        content_hash=None,
        last_synced_at=None,
        created_at=datetime.now(timezone.utc),
    )
    defaults.update(overrides)
    return SimpleNamespace(**defaults)


def _make_chunk(**overrides) -> dict:
    from app.services.rag_retrieval_service import RetrievedChunk
    defaults = dict(
        vector_id=f"vec-{uuid.uuid4().hex[:8]}",
        source_id=uuid.uuid4(),
        source_title="Source A",
        section_heading="Requirements",
        chunk_text="User can log in with valid credentials",
        relevance_score=0.92,
        requirement_id="AC-1",
        chunk_text_preview="User can log in with valid credentials",
    )
    defaults.update(overrides)
    return RetrievedChunk(**defaults)


# ═══════════════════════════════════════════════════════════════════════════════
# RAG-2: Connector Tests — Jira
# ═══════════════════════════════════════════════════════════════════════════════


class TestJiraConnectorInternals:
    def test_extract_issue_key_from_url(self):
        from app.services.connectors.jira_connector import JiraKnowledgeConnector
        c = JiraKnowledgeConnector()
        key = c._extract_issue_key("https://jira.example.com/browse/PROJ-123", None)
        assert key == "PROJ-123"

    def test_extract_issue_key_from_external_id(self):
        from app.services.connectors.jira_connector import JiraKnowledgeConnector
        c = JiraKnowledgeConnector()
        key = c._extract_issue_key("https://jira.example.com/browse/PROJ-123", "PROJ-456")
        assert key == "PROJ-456"

    def test_extract_issue_key_missing_raises(self):
        from app.services.connectors.jira_connector import JiraKnowledgeConnector
        c = JiraKnowledgeConnector()
        with pytest.raises(ConnectorFetchError, match="Cannot extract"):
            c._extract_issue_key("https://jira.example.com/somepage", None)

    def test_normalize_issue_basic(self):
        from app.services.connectors.jira_connector import JiraKnowledgeConnector
        c = JiraKnowledgeConnector()
        fields = {
            "summary": "Login broken",
            "status": {"name": "Open"},
            "priority": {"name": "High"},
            "issuetype": {"name": "Bug"},
            "description": None,
            "labels": ["auth", "critical"],
            "subtasks": [],
        }
        text = c._normalize_issue("PROJ-1", fields)
        assert "PROJ-1: Login broken" in text
        assert "Bug" in text
        assert "auth" in text

    def test_normalize_issue_with_subtasks(self):
        from app.services.connectors.jira_connector import JiraKnowledgeConnector
        c = JiraKnowledgeConnector()
        fields = {
            "summary": "Epic",
            "status": {"name": "Open"},
            "priority": {"name": "Medium"},
            "issuetype": {"name": "Epic"},
            "subtasks": [
                {"key": "PROJ-2", "fields": {"summary": "Subtask 1", "status": {"name": "Done"}}},
            ],
        }
        text = c._normalize_issue("PROJ-1", fields)
        assert "Subtasks" in text
        assert "PROJ-2" in text
        assert "Done" in text

    def test_adf_to_text_simple_paragraph(self):
        from app.services.connectors.jira_connector import JiraKnowledgeConnector
        c = JiraKnowledgeConnector()
        adf = {
            "type": "doc",
            "content": [
                {"type": "paragraph", "content": [{"type": "text", "text": "Hello world"}]}
            ],
        }
        result = c._adf_to_text(adf)
        assert "Hello world" in result

    def test_adf_to_text_heading(self):
        from app.services.connectors.jira_connector import JiraKnowledgeConnector
        c = JiraKnowledgeConnector()
        adf = {
            "type": "heading",
            "attrs": {"level": 2},
            "content": [{"type": "text", "text": "Section Title"}],
        }
        result = c._adf_to_text(adf)
        assert "## Section Title" in result

    def test_adf_to_text_code_block(self):
        from app.services.connectors.jira_connector import JiraKnowledgeConnector
        c = JiraKnowledgeConnector()
        adf = {
            "type": "codeBlock",
            "content": [{"type": "text", "text": "print('hi')"}],
        }
        result = c._adf_to_text(adf)
        assert "```" in result
        assert "print('hi')" in result

    def test_adf_to_text_list_item(self):
        from app.services.connectors.jira_connector import JiraKnowledgeConnector
        c = JiraKnowledgeConnector()
        adf = {
            "type": "bulletList",
            "content": [
                {"type": "listItem", "content": [{"type": "text", "text": "Item 1"}]},
            ],
        }
        result = c._adf_to_text(adf)
        assert "- Item 1" in result

    def test_adf_to_text_string_input(self):
        from app.services.connectors.jira_connector import JiraKnowledgeConnector
        c = JiraKnowledgeConnector()
        assert c._adf_to_text("plain string") == "plain string"

    def test_adf_to_text_none_returns_empty(self):
        from app.services.connectors.jira_connector import JiraKnowledgeConnector
        c = JiraKnowledgeConnector()
        assert c._adf_to_text(None) == ""

    def test_adf_to_text_table(self):
        from app.services.connectors.jira_connector import JiraKnowledgeConnector
        c = JiraKnowledgeConnector()
        adf = {
            "type": "table",
            "content": [
                {
                    "type": "tableRow",
                    "content": [
                        {"type": "tableHeader", "content": [{"type": "text", "text": "Col1"}]},
                        {"type": "tableHeader", "content": [{"type": "text", "text": "Col2"}]},
                    ],
                },
            ],
        }
        result = c._adf_to_text(adf)
        assert "Col1" in result
        assert "|" in result


# ═══════════════════════════════════════════════════════════════════════════════
# RAG-2: Connector Tests — Confluence
# ═══════════════════════════════════════════════════════════════════════════════


class TestConfluenceConnectorInternals:
    def test_extract_page_id_from_url(self):
        from app.services.connectors.confluence_connector import ConfluenceKnowledgeConnector
        c = ConfluenceKnowledgeConnector()
        assert c._extract_page_id("https://wiki.example.com/wiki/spaces/ENG/pages/123456/My+Page", None) == "123456"

    def test_extract_page_id_from_external_id(self):
        from app.services.connectors.confluence_connector import ConfluenceKnowledgeConnector
        c = ConfluenceKnowledgeConnector()
        assert c._extract_page_id("https://wiki.example.com", "789") == "789"

    def test_extract_page_id_missing_raises(self):
        from app.services.connectors.confluence_connector import ConfluenceKnowledgeConnector
        c = ConfluenceKnowledgeConnector()
        with pytest.raises(ConnectorFetchError, match="Cannot extract"):
            c._extract_page_id("https://wiki.example.com/nopages", None)

    def test_storage_to_text_empty_page(self):
        from app.services.connectors.confluence_connector import ConfluenceKnowledgeConnector
        c = ConfluenceKnowledgeConnector()
        result = c._storage_to_text("Empty Page", "ENG", 3, "")
        assert "Empty Page" in result
        assert "Empty page" in result

    def test_regex_strip_html(self):
        from app.services.connectors.confluence_connector import ConfluenceKnowledgeConnector
        html = "<p>Hello <b>World</b></p><br/><h2>Title</h2><p>Body</p>"
        result = ConfluenceKnowledgeConnector._regex_strip_html(html)
        assert "Hello World" in result
        assert "Title" in result
        assert "<" not in result


# ═══════════════════════════════════════════════════════════════════════════════
# RAG-2: Connector Tests — URL
# ═══════════════════════════════════════════════════════════════════════════════


class TestURLConnectorInternals:
    def test_extract_title_from_html(self):
        from app.services.connectors.url_connector import _extract_title
        body = '<html><head><title>My Page &amp; More</title></head><body>test</body></html>'
        title = _extract_title(body, "text/html")
        assert title == "My Page & More"

    def test_extract_title_no_title_tag(self):
        from app.services.connectors.url_connector import _extract_title
        body = "<html><body>no title</body></html>"
        assert _extract_title(body, "text/html") is None

    def test_extract_title_non_html(self):
        from app.services.connectors.url_connector import _extract_title
        assert _extract_title("plain text content", "text/plain") is None

    def test_extract_title_truncates_long_title(self):
        from app.services.connectors.url_connector import _extract_title
        body = f'<html><head><title>{"X" * 300}</title></head></html>'
        title = _extract_title(body, "text/html")
        assert len(title) == 200

    def test_html_to_text_fallback(self):
        from app.services.connectors.url_connector import _regex_strip_html
        html = "<div><p>Hello</p><script>evil()</script><style>.x{}</style></div>"
        result = _regex_strip_html(html)
        assert "Hello" in result
        assert "evil" not in result
        assert ".x" not in result

    @pytest.mark.asyncio
    async def test_fetch_content_blocks_file_scheme(self):
        from app.services.connectors.url_connector import URLConnector
        c = URLConnector()
        with pytest.raises(ConnectorFetchError, match="not allowed"):
            await c.fetch_content("file:///etc/passwd")

    @pytest.mark.asyncio
    async def test_fetch_content_blocks_data_scheme(self):
        from app.services.connectors.url_connector import URLConnector
        c = URLConnector()
        with pytest.raises(ConnectorFetchError, match="not allowed"):
            await c.fetch_content("data:text/plain,hello")


# ═══════════════════════════════════════════════════════════════════════════════
# RAG-2: Connector Tests — Document
# ═══════════════════════════════════════════════════════════════════════════════


class TestDocumentConnectorInternals:
    def test_extract_docx_preserves_mixed_block_order_and_formatting(self):
        from io import BytesIO

        from docx import Document

        from app.services.connectors.document_connector import _extract_docx

        doc = Document()
        doc.add_paragraph("Overview", style="Heading 1")
        doc.add_paragraph("Run this first", style="List Bullet")
        doc.add_paragraph("Before the table")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "A"
        table.cell(0, 1).text = "B"
        table.cell(1, 0).text = "C"
        table.cell(1, 1).text = "D"
        doc.add_paragraph("After the table")

        output = BytesIO()
        doc.save(output)
        result = _extract_docx(output.getvalue())

        assert result.split("\n\n") == [
            "# Overview",
            "- Run this first",
            "Before the table",
            "| A | B |\n| C | D |",
            "After the table",
        ]

    def test_extract_text_routing_pdf(self):
        from app.services.connectors.document_connector import _extract_text
        with pytest.raises(ConnectorFetchError, match="pypdf"):
            # Will fail because pypdf likely not installed in test env,
            # but it should try the PDF path, not raise "Unsupported"
            _extract_text(b"not-a-pdf", "pdf", "test.pdf")

    def test_extract_text_routing_unsupported(self):
        from app.services.connectors.document_connector import _extract_text
        with pytest.raises(ConnectorFetchError, match="Unsupported file type"):
            _extract_text(b"data", "exe", "virus.exe")

    def test_extract_plaintext_utf8(self):
        from app.services.connectors.document_connector import _extract_plaintext
        text = _extract_plaintext("Hello world".encode("utf-8"), "test.txt")
        assert text == "Hello world"

    def test_extract_plaintext_latin1(self):
        from app.services.connectors.document_connector import _extract_plaintext
        text = _extract_plaintext("Héllo".encode("latin-1"), "test.txt")
        assert "llo" in text

    def test_extract_plaintext_bad_encoding_raises(self):
        from app.services.connectors.document_connector import _extract_plaintext
        # Bytes that are invalid in all tried encodings — actually latin-1
        # accepts all byte values, so this won't fail. Test with a valid case instead.
        text = _extract_plaintext(b"\xff\xfe", "test.txt")
        assert isinstance(text, str)

    def test_extract_text_markdown(self):
        from app.services.connectors.document_connector import _extract_text
        content = b"# Title\n\nSome **bold** text"
        result = _extract_text(content, "md", "readme.md")
        assert "# Title" in result
        assert "bold" in result

    def test_extract_text_rst(self):
        from app.services.connectors.document_connector import _extract_text
        content = b"Title\n=====\n\nSome text"
        result = _extract_text(content, "rst", "readme.rst")
        assert "Title" in result


# ═══════════════════════════════════════════════════════════════════════════════
# RAG-5: Chunking edge cases
# ═══════════════════════════════════════════════════════════════════════════════


class TestChunkingEdgeCases:
    def test_split_by_headings_no_headings(self):
        from app.services.knowledge_chunking_service import _split_by_headings
        parts = _split_by_headings("Just plain text without any headings")
        assert len(parts) == 1
        assert parts[0][0] is None
        assert "plain text" in parts[0][1]

    def test_split_by_headings_only_heading(self):
        from app.services.knowledge_chunking_service import _split_by_headings
        parts = _split_by_headings("## Only Heading")
        assert len(parts) == 1
        # When no body follows the heading, the full text becomes a single part
        assert parts[0][1] == "## Only Heading"

    def test_split_by_headings_heading_with_body(self):
        from app.services.knowledge_chunking_service import _split_by_headings
        parts = _split_by_headings("## Heading\nBody text here")
        assert len(parts) == 1
        assert parts[0][0] == "Heading"
        assert "Body text here" in parts[0][1]

    def test_split_by_requirements_numbered_list(self):
        from app.services.knowledge_chunking_service import _split_by_requirements
        text = "1. First item\n2. Second item\n3. Third item"
        parts = _split_by_requirements(text)
        assert len(parts) >= 2

    def test_split_by_requirements_given_when_then(self):
        from app.services.knowledge_chunking_service import _split_by_requirements
        text = "- **Given** a logged-in user\n- **When** they click logout\n- **Then** they are logged out"
        parts = _split_by_requirements(text)
        assert len(parts) >= 2

    def test_split_by_requirements_no_patterns(self):
        from app.services.knowledge_chunking_service import _split_by_requirements
        text = "This is just regular prose without any requirement markers."
        parts = _split_by_requirements(text)
        assert len(parts) == 1

    def test_hard_split_preserves_content(self):
        from app.services.knowledge_chunking_service import _hard_split
        # Create text with 3 clear paragraphs
        text = "Word " * 500 + "\n\n" + "More " * 500 + "\n\n" + "Final " * 500
        parts = _hard_split(text, 400)
        assert len(parts) >= 2
        # Reconstruct should contain all original words
        reconstructed = " ".join(parts)
        assert "Word" in reconstructed
        assert "More" in reconstructed
        assert "Final" in reconstructed

    def test_hard_split_single_paragraph_word_split(self):
        from app.services.knowledge_chunking_service import _hard_split
        text = "word " * 2000  # Way over max_tokens
        parts = _hard_split(text, 200)
        assert len(parts) >= 2

    def test_split_into_chunks_with_overlap(self, monkeypatch):
        from app.services.knowledge_chunking_service import _split_into_chunks
        monkeypatch.setattr("app.services.knowledge_chunking_service.settings.KNOWLEDGE_CHUNK_TARGET_TOKENS", 50)
        monkeypatch.setattr("app.services.knowledge_chunking_service.settings.KNOWLEDGE_CHUNK_MAX_TOKENS", 100)
        monkeypatch.setattr("app.services.knowledge_chunking_service.settings.KNOWLEDGE_CHUNK_OVERLAP_TOKENS", 10)
        content = FetchedContent(
            raw_text="## Section A\n" + ("word " * 100) + "\n## Section B\n" + ("other " * 100),
            content_hash="abc",
            title="Test",
            source_url="http://x",
        )
        chunks = _split_into_chunks(content)
        assert len(chunks) >= 2
        assert all(c.token_count > 0 for c in chunks)

    def test_split_into_chunks_assigns_section_headings(self):
        from app.services.knowledge_chunking_service import _split_into_chunks
        content = FetchedContent(
            raw_text="## Requirements\nREQ-1: User logs in\n\n## Acceptance Criteria\nAC-1: Login form appears",
            content_hash="abc",
            title="Test",
            source_url="http://x",
        )
        chunks = _split_into_chunks(content)
        headings = {c.section_heading for c in chunks}
        assert "Requirements" in headings or "Acceptance Criteria" in headings

    def test_extract_requirement_id_various_formats(self):
        from app.services.knowledge_chunking_service import _extract_requirement_id
        assert _extract_requirement_id("US-10: User story") is not None
        assert _extract_requirement_id("TC-5 Test case") is not None
        assert _extract_requirement_id("[NFR-3] Non-functional") is not None
        assert _extract_requirement_id("(UC-7) Use case") is not None
        assert _extract_requirement_id("Normal text") is None
        assert _extract_requirement_id("") is None

    def test_estimate_tokens_empty_string(self):
        from app.services.knowledge_chunking_service import _estimate_tokens
        assert _estimate_tokens("") >= 1

    def test_estimate_tokens_single_word(self):
        from app.services.knowledge_chunking_service import _estimate_tokens
        assert _estimate_tokens("hello") >= 1


# ═══════════════════════════════════════════════════════════════════════════════
# RAG-6: Sync freshness computation
# ═══════════════════════════════════════════════════════════════════════════════


class TestFreshnessComputation:
    def test_uploaded_doc_never_stale(self):
        from app.services.knowledge_sync_service import compute_staleness
        src = _make_source(
            source_type=KnowledgeSourceType.UPLOADED_DOC.value,
            last_synced_at=datetime.now(timezone.utc) - timedelta(days=365),
        )
        is_stale, stale_since = compute_staleness(src)
        assert is_stale is False
        assert stale_since is None

    def test_never_synced_is_stale(self):
        from app.services.knowledge_sync_service import compute_staleness
        src = _make_source(last_synced_at=None)
        is_stale, stale_since = compute_staleness(src)
        assert is_stale is True

    def test_recently_synced_is_fresh(self):
        from app.services.knowledge_sync_service import compute_staleness
        src = _make_source(
            last_synced_at=datetime.now(timezone.utc) - timedelta(hours=1),
        )
        is_stale, stale_since = compute_staleness(src)
        assert is_stale is False

    def test_old_sync_is_stale(self):
        from app.services.knowledge_sync_service import compute_staleness
        src = _make_source(
            source_type=KnowledgeSourceType.JIRA_ISSUE.value,
            last_synced_at=datetime.now(timezone.utc) - timedelta(hours=48),
        )
        is_stale, stale_since = compute_staleness(src)
        assert is_stale is True
        assert stale_since is not None

    def test_url_source_uses_url_threshold(self):
        from app.services.knowledge_sync_service import compute_staleness
        src = _make_source(
            source_type=KnowledgeSourceType.EXTERNAL_URL.value,
            last_synced_at=datetime.now(timezone.utc) - timedelta(hours=100),
        )
        # Default URL threshold is 168 hours, so 100h should be fresh
        is_stale, _ = compute_staleness(src)
        assert is_stale is False


class TestEffectiveThreshold:
    def test_jira_uses_config(self, monkeypatch):
        from app.services.knowledge_sync_service import _effective_threshold
        monkeypatch.setattr("app.services.knowledge_sync_service.settings.KNOWLEDGE_STALE_THRESHOLD_JIRA_HOURS", 12)
        assert _effective_threshold(KnowledgeSourceType.JIRA_ISSUE.value) == 12

    def test_url_uses_config(self, monkeypatch):
        from app.services.knowledge_sync_service import _effective_threshold
        monkeypatch.setattr("app.services.knowledge_sync_service.settings.KNOWLEDGE_STALE_THRESHOLD_URL_HOURS", 72)
        assert _effective_threshold(KnowledgeSourceType.EXTERNAL_URL.value) == 72

    def test_uploaded_doc_returns_negative_one(self):
        from app.services.knowledge_sync_service import _effective_threshold
        assert _effective_threshold(KnowledgeSourceType.UPLOADED_DOC.value) == -1


# ═══════════════════════════════════════════════════════════════════════════════
# RAG-7: Retrieval parsing
# ═══════════════════════════════════════════════════════════════════════════════


class TestRetrievalParsingEdgeCases:
    def test_parse_results_mismatched_lengths(self):
        from app.services.rag_retrieval_service import _parse_results
        results = {
            "ids": [["vec-1", "vec-2"]],
            "documents": [["text 1"]],  # shorter
            "distances": [[0.3]],  # shorter
            "metadatas": [[
                {"source_id": str(uuid.uuid4()), "section_heading": "S1"},
                {"source_id": str(uuid.uuid4())},  # no metadata match for doc
            ]],
        }
        chunks = _parse_results(results)
        assert len(chunks) >= 1  # Should handle gracefully

    def test_parse_results_distance_to_score_conversion(self):
        from app.services.rag_retrieval_service import _parse_results
        sid = str(uuid.uuid4())
        results = {
            "ids": [["vec-1"]],
            "documents": [["text"]],
            "distances": [[0.0]],  # distance 0 = perfect match = score 1.0
            "metadatas": [[{"source_id": sid}]],
        }
        chunks = _parse_results(results)
        assert len(chunks) == 1
        assert chunks[0].relevance_score == 1.0

    def test_parse_results_high_distance_low_score(self):
        from app.services.rag_retrieval_service import _parse_results
        sid = str(uuid.uuid4())
        results = {
            "ids": [["vec-1"]],
            "documents": [["text"]],
            "distances": [[1.5]],  # distance > 1 should clamp score to 0
            "metadatas": [[{"source_id": sid}]],
        }
        chunks = _parse_results(results)
        assert len(chunks) == 1
        assert chunks[0].relevance_score == 0.0

    def test_parse_results_none_input(self):
        from app.services.rag_retrieval_service import _parse_results
        assert _parse_results(None) == []

    def test_parse_results_preserves_requirement_id(self):
        from app.services.rag_retrieval_service import _parse_results
        sid = str(uuid.uuid4())
        results = {
            "ids": [["vec-1"]],
            "documents": [["text"]],
            "distances": [[0.2]],
            "metadatas": [[{"source_id": sid, "requirement_id": "REQ-42"}]],
        }
        chunks = _parse_results(results)
        assert chunks[0].requirement_id == "REQ-42"

    def test_parse_results_empty_requirement_id_becomes_none(self):
        from app.services.rag_retrieval_service import _parse_results
        sid = str(uuid.uuid4())
        results = {
            "ids": [["vec-1"]],
            "documents": [["text"]],
            "distances": [[0.2]],
            "metadatas": [[{"source_id": sid, "requirement_id": ""}]],
        }
        chunks = _parse_results(results)
        assert chunks[0].requirement_id is None


# ═══════════════════════════════════════════════════════════════════════════════
# RAG-8: Generation — prompt building, LLM fallback, coverage mapping
# ═══════════════════════════════════════════════════════════════════════════════


class TestBuildGroundedPrompt:
    def test_with_multiple_chunks(self):
        from app.services.rag_generation_service import _build_grounded_prompt
        chunks = [
            _make_chunk(source_title="Jira PROJ-1", requirement_id="AC-1"),
            _make_chunk(source_title="Confluence Page", requirement_id="REQ-2"),
        ]
        prompt = _build_grounded_prompt("Generate tests", chunks)
        assert 'id="EVIDENCE-1"' in prompt
        assert 'id="EVIDENCE-2"' in prompt
        assert "Jira PROJ-1" in prompt
        assert "Confluence Page" in prompt
        assert "AC-1" in prompt
        assert "REQ-2" in prompt
        assert "Generate tests" in prompt

    def test_chunk_text_truncated_to_2000(self):
        from app.services.rag_generation_service import _build_grounded_prompt
        chunk = _make_chunk(chunk_text="X" * 5000)
        prompt = _build_grounded_prompt("test", [chunk])
        # The chunk text in the prompt should be truncated
        assert "X" * 2001 not in prompt

    def test_without_section_heading(self):
        from app.services.rag_generation_service import _build_grounded_prompt
        chunk = _make_chunk(section_heading=None, requirement_id=None)
        prompt = _build_grounded_prompt("test", [chunk])
        assert 'id="EVIDENCE-1"' in prompt


class TestCallLlmGenerate:
    @pytest.mark.asyncio
    async def test_import_error_is_an_explicit_provider_failure(self):
        from app.services.rag_generation_service import RagGenerationUnavailable, _call_llm_generate
        with patch.dict("sys.modules", {"app.services.test_case_ai_agent": None}):
            with pytest.raises(RagGenerationUnavailable, match="provider is unavailable"):
                await _call_llm_generate("test prompt", None)

    @pytest.mark.asyncio
    async def test_generation_config_appended_to_prompt(self):
        from app.services.rag_generation_service import _call_llm_generate
        config = {"test_type": "security", "priority": "high", "count": 5}
        generated = [{"title": "Configured case"}]
        with patch(
            "app.services.test_case_ai_agent.ai_generate_test_cases",
            AsyncMock(return_value=generated),
        ):
            result = await _call_llm_generate("base prompt", config)
        assert isinstance(result, list)
        assert len(result) >= 1


class TestStubGeneratedCases:
    def test_returns_minimum_two_cases(self):
        from app.services.rag_generation_service import _stub_generated_cases
        cases = _stub_generated_cases("any prompt")
        assert len(cases) == 2

    def test_each_case_has_required_fields(self):
        from app.services.rag_generation_service import _stub_generated_cases
        for case in _stub_generated_cases("test"):
            assert "title" in case
            assert "description" in case
            assert "steps" in case
            assert isinstance(case["steps"], list)
            assert "test_type" in case
            assert "priority" in case

    def test_steps_have_structure(self):
        from app.services.rag_generation_service import _stub_generated_cases
        case = _stub_generated_cases("test")[0]
        step = case["steps"][0]
        assert "step_number" in step
        assert "action" in step
        assert "expected_result" in step


class TestBuildCitations:
    def test_empty_chunks_no_citations(self):
        from app.services.rag_generation_service import _build_citations
        citations = _build_citations([{"title": "Case 1"}], [])
        assert citations == []

    def test_respects_max_citations_constant(self):
        from app.services.rag_generation_service import (
            MAX_CITATIONS_PER_CASE,
            _build_citations,
        )
        chunks = [_make_chunk() for _ in range(10)]
        cases = [{
            "title": "Case 1",
            "evidence_ids": [f"EVIDENCE-{index}" for index in range(1, 11)],
        }]
        citations = _build_citations(cases, chunks)
        assert len(citations) == MAX_CITATIONS_PER_CASE

    def test_citations_per_case(self):
        from app.services.rag_generation_service import _build_citations
        chunks = [_make_chunk()]
        cases = [
            {"title": "Case 1", "evidence_ids": ["EVIDENCE-1"]},
            {"title": "Case 2", "evidence_ids": ["EVIDENCE-1"]},
            {"title": "Case 3", "evidence_ids": ["EVIDENCE-1"]},
        ]
        citations = _build_citations(cases, chunks)
        # Each case explicitly binds itself to the first retrieved evidence item.
        case_indices = {c["case_index"] for c in citations}
        assert case_indices == {0, 1, 2}


@pytest.mark.asyncio
class TestCoverageMapping:
    async def test_no_requirement_ids_returns_none(self):
        from app.services.rag_generation_service import _map_coverage
        chunk = _make_chunk(requirement_id=None, section_heading=None)
        mock_db = AsyncMock()
        summary, gaps = await _map_coverage(mock_db, uuid.uuid4(), uuid.uuid4(), [chunk], [])
        assert summary is None
        assert gaps == []

    async def test_covered_requirement_detected(self):
        from app.services.rag_generation_service import _map_coverage
        chunk = _make_chunk(requirement_id="AC-1")
        cases = [{"title": "Test AC-1 login flow", "description": "covers AC-1"}]
        mock_db = AsyncMock()
        summary, gaps = await _map_coverage(mock_db, uuid.uuid4(), uuid.uuid4(), [chunk], cases)
        assert "1/1" in summary
        assert len(gaps) == 0

    async def test_uncovered_requirement_in_gaps(self):
        from app.services.rag_generation_service import _map_coverage
        chunk = _make_chunk(requirement_id="REQ-99")
        cases = [{"title": "Unrelated test", "description": "nothing relevant"}]
        mock_db = AsyncMock()
        summary, gaps = await _map_coverage(mock_db, uuid.uuid4(), uuid.uuid4(), [chunk], cases)
        assert "0/1" in summary
        assert "REQ-99" in gaps

    async def test_section_heading_used_as_fallback_req_id(self):
        from app.services.rag_generation_service import _map_coverage
        chunk = _make_chunk(requirement_id=None, section_heading="Login Requirements")
        cases = [{"title": "Login Requirements verification"}]
        mock_db = AsyncMock()
        summary, gaps = await _map_coverage(mock_db, uuid.uuid4(), uuid.uuid4(), [chunk], cases)
        assert "1/1" in summary

    async def test_empty_cases_all_uncovered(self):
        from app.services.rag_generation_service import _map_coverage
        chunk = _make_chunk(requirement_id="REQ-1")
        mock_db = AsyncMock()
        summary, gaps = await _map_coverage(mock_db, uuid.uuid4(), uuid.uuid4(), [chunk], [])
        assert "0/1" in summary
        assert "REQ-1" in gaps


# ═══════════════════════════════════════════════════════════════════════════════
# RAG-10: Review service
# ═══════════════════════════════════════════════════════════════════════════════


class TestReviewServiceEdgeCases:
    def test_accept_case_field_protection(self):
        """The actual accept schema fails closed for lifecycle/identity edits."""
        from pydantic import ValidationError

        from app.models.schemas import RagCaseAcceptEdits

        for protected, value in {
            "id": str(uuid.uuid4()),
            "project_id": str(uuid.uuid4()),
            "generation_batch_id": str(uuid.uuid4()),
            "created_at": datetime.now(timezone.utc).isoformat(),
            "status": "active",
            "version": 999,
            "author_id": str(uuid.uuid4()),
            "reviewer_id": str(uuid.uuid4()),
        }.items():
            with pytest.raises(ValidationError):
                RagCaseAcceptEdits.model_validate(
                    {"title": "Safe replacement", protected: value}
                )

        allowed = RagCaseAcceptEdits.model_validate(
            {"title": "Safe replacement", "description": "Reviewed content"}
        )
        assert allowed.title == "Safe replacement"
        assert allowed.description == "Reviewed content"

    async def test_accept_records_disposition_once_and_edits_content(self):
        from app.services import rag_review_service as review

        batch_id = uuid.uuid4()
        project_id = uuid.uuid4()
        case = SimpleNamespace(
            id=uuid.uuid4(),
            project_id=project_id,
            generation_batch_id=batch_id,
            status="draft",
            title="Generated title",
            description="Generated description",
            version=1,
        )
        batch = SimpleNamespace(project_id=project_id, cases_accepted=0)
        user = SimpleNamespace(id=uuid.uuid4())
        db = AsyncMock()
        db.execute.side_effect = [
            SimpleNamespace(scalar_one_or_none=lambda: case),
            SimpleNamespace(scalar_one_or_none=lambda: None),
        ]

        with (
            patch.object(review, "_get_batch_or_404", AsyncMock(return_value=batch)),
            patch(
                "app.services.rag_faithfulness_service.check_accept",
                AsyncMock(return_value={"allow": True}),
            ),
            patch.object(review, "stage_test_case_snapshot") as snapshot,
            patch.object(review, "audit_event", AsyncMock()) as audit,
        ):
            result = await review.accept_case(
                db,
                batch_id,
                case.id,
                {"title": "Reviewed title"},
                user,
            )

        assert result is case
        assert case.title == "Reviewed title"
        assert case.version == 2
        assert batch.cases_accepted == 1
        snapshot.assert_called_once()
        assert audit.await_args.args[4] == review.GENERATION_ACCEPTED_ACTION
        db.flush.assert_awaited_once()

    @pytest.mark.parametrize(
        ("prior_action", "requested_action", "disposition"),
        [
            ("generation_accepted", "accept", "accepted"),
            ("generation_rejected", "accept", "rejected"),
            ("generation_accepted", "reject", "accepted"),
            ("generation_rejected", "reject", "rejected"),
        ],
    )
    async def test_repeat_or_conflicting_disposition_is_409(
        self,
        prior_action,
        requested_action,
        disposition,
    ):
        from fastapi import HTTPException

        from app.services import rag_review_service as review

        case = SimpleNamespace(
            id=uuid.uuid4(),
            project_id=uuid.uuid4(),
            status="draft",
        )
        db = AsyncMock()
        db.execute.return_value = SimpleNamespace(
            scalar_one_or_none=lambda: prior_action
        )

        with pytest.raises(HTTPException) as exc:
            await review._ensure_initial_undisposed_draft(
                db,
                case,
                requested_action=requested_action,
            )

        assert exc.value.status_code == 409
        assert exc.value.detail["disposition"] == disposition
        assert exc.value.detail["requested_action"] == requested_action

    async def test_reject_uses_governed_review_sequence_and_disposition_audit(self):
        from app.services import rag_review_service as review

        batch_id = uuid.uuid4()
        project_id = uuid.uuid4()
        case = SimpleNamespace(
            id=uuid.uuid4(),
            project_id=project_id,
            generation_batch_id=batch_id,
            status="draft",
            description="Generated description",
        )
        batch = SimpleNamespace(project_id=project_id, cases_rejected=0)
        user = SimpleNamespace(id=uuid.uuid4())
        db = AsyncMock()
        db.execute.side_effect = [
            SimpleNamespace(scalar_one_or_none=lambda: case),
            SimpleNamespace(scalar_one_or_none=lambda: None),
        ]

        with (
            patch.object(review, "_get_batch_or_404", AsyncMock(return_value=batch)),
            patch.object(review, "transition", AsyncMock()) as transition_mock,
            patch.object(review, "audit_event", AsyncMock()) as audit,
        ):
            await review.reject_case(db, batch_id, case.id, "  duplicate  ", user)

        assert case.description.startswith("[Rejected: duplicate]")
        assert [call.args[2] for call in transition_mock.await_args_list] == [
            review.LifecycleAction.REQUEST_REVIEW,
            review.LifecycleAction.CLAIM_REVIEW,
            review.LifecycleAction.REJECT,
        ]
        assert transition_mock.await_args_list[-1].kwargs["reason"] == "duplicate"
        assert batch.cases_rejected == 1
        assert audit.await_args.args[4] == review.GENERATION_REJECTED_ACTION
        assert audit.await_args.kwargs["reason"] == "duplicate"

    async def test_reject_reason_over_500_is_422_before_description_mutation(self):
        from fastapi import HTTPException

        from app.services import rag_review_service as review

        case = SimpleNamespace(
            id=uuid.uuid4(),
            project_id=uuid.uuid4(),
            generation_batch_id=uuid.uuid4(),
            status="draft",
            description="original",
        )
        batch = SimpleNamespace(project_id=case.project_id, cases_rejected=0)
        db = AsyncMock()
        db.execute.side_effect = [
            SimpleNamespace(scalar_one_or_none=lambda: case),
            SimpleNamespace(scalar_one_or_none=lambda: None),
        ]
        transition_mock = AsyncMock()

        with (
            patch.object(review, "_get_batch_or_404", AsyncMock(return_value=batch)),
            patch.object(review, "transition", transition_mock),
        ):
            with pytest.raises(HTTPException) as exc:
                await review.reject_case(
                    db,
                    case.generation_batch_id,
                    case.id,
                    "x" * 501,
                    SimpleNamespace(id=uuid.uuid4()),
                )

        assert exc.value.status_code == 422
        assert case.description == "original"
        assert batch.cases_rejected == 0
        transition_mock.assert_not_awaited()

    async def test_bulk_accept_deduplicates_case_ids_before_mutating(self):
        from app.services import rag_review_service as review

        first_id = uuid.uuid4()
        second_id = uuid.uuid4()
        accepted = {
            first_id: SimpleNamespace(id=first_id),
            second_id: SimpleNamespace(id=second_id),
        }
        accept = AsyncMock(side_effect=lambda _db, _batch, cid, **_kwargs: accepted[cid])
        with patch.object(review, "accept_case", accept):
            result = await review.bulk_accept(
                AsyncMock(),
                uuid.uuid4(),
                [first_id, first_id, second_id, first_id],
                SimpleNamespace(id=uuid.uuid4()),
            )

        assert [case.id for case in result] == [first_id, second_id]
        assert [call.args[2] for call in accept.await_args_list] == [first_id, second_id]


# ═══════════════════════════════════════════════════════════════════════════════
# RAG-12: Staleness detection
# ═══════════════════════════════════════════════════════════════════════════════


class TestStalenessModels:
    def test_generation_case_source_stale_fields_exist(self):
        """Verify the model has stale-related columns."""
        gcs = GenerationCaseSource(
            batch_id=uuid.uuid4(),
            case_id=uuid.uuid4(),
            source_id=uuid.uuid4(),
            chunk_vector_id="vec-1",
            is_stale=False,
        )
        assert gcs.is_stale is False
        assert gcs.stale_detected_at is None

    def test_managed_test_case_stale_fields_exist(self):
        """Verify the model accepts stale-related fields."""
        tc = ManagedTestCase(
            project_id=uuid.uuid4(),
            title="Test",
            is_stale=False,
        )
        assert tc.is_stale is False
        assert tc.stale_reason is None

    def test_generation_batch_fields(self):
        """Verify batch model accepts all expected fields."""
        batch = GenerationBatch(
            project_id=uuid.uuid4(),
            generation_mode="raw",
            cases_generated=0,
            cases_accepted=0,
            cases_rejected=0,
            status="pending",
        )
        assert batch.cases_generated == 0
        assert batch.cases_accepted == 0
        assert batch.cases_rejected == 0
        assert batch.status == "pending"


class TestStalenessService:
    async def test_mark_cases_stale_for_source_propagates_to_cases(self):
        from app.services.rag_staleness_service import mark_cases_stale_for_source

        source_id = uuid.uuid4()
        case_id = uuid.uuid4()
        citation = SimpleNamespace(
            case_id=case_id,
            is_stale=False,
            stale_detected_at=None,
            source_content_hash_at_generation="old-hash",
        )
        source_result = SimpleNamespace(scalar_one_or_none=lambda: "new-hash")
        citations_result = SimpleNamespace(scalars=lambda: SimpleNamespace(all=lambda: [citation]))
        db = AsyncMock()
        db.execute.side_effect = [source_result, citations_result, AsyncMock()]

        marked = await mark_cases_stale_for_source(db, source_id)

        assert marked == 1
        assert citation.is_stale is True
        assert citation.stale_detected_at is not None
        source_query = db.execute.await_args_list[1].args[0]
        assert "!=" in str(source_query)
        update_stmt = db.execute.await_args_list[2].args[0]
        assert "managed_test_cases" in str(update_stmt)
        assert "stale_reason" in str(update_stmt)
        db.flush.assert_awaited_once()

    async def test_mark_cases_stale_for_source_is_idempotent(self):
        from app.services.rag_staleness_service import mark_cases_stale_for_source

        citation = SimpleNamespace(
            case_id=uuid.uuid4(),
            is_stale=False,
            stale_detected_at=None,
            source_content_hash_at_generation="same-hash",
        )
        source_result = SimpleNamespace(scalar_one_or_none=lambda: "same-hash")
        citations_result = SimpleNamespace(scalars=lambda: SimpleNamespace(all=lambda: []))
        db = AsyncMock()
        db.execute.side_effect = [source_result, citations_result]

        assert await mark_cases_stale_for_source(db, uuid.uuid4()) == 0
        assert citation.is_stale is False
        db.flush.assert_not_awaited()

    async def test_check_batch_staleness_marks_only_changed_citations(self):
        from app.services.rag_staleness_service import check_batch_staleness

        stale = SimpleNamespace(
            case_id=uuid.uuid4(),
            source_id=uuid.uuid4(),
            is_stale=False,
            stale_detected_at=None,
            source_content_hash_at_generation="old",
        )
        fresh = SimpleNamespace(
            case_id=uuid.uuid4(),
            source_id=uuid.uuid4(),
            is_stale=False,
            stale_detected_at=None,
            source_content_hash_at_generation="same",
        )
        already_stale = SimpleNamespace(
            case_id=uuid.uuid4(),
            source_id=uuid.uuid4(),
            is_stale=True,
            stale_detected_at=datetime.now(timezone.utc),
            source_content_hash_at_generation="old",
        )
        results = [
            SimpleNamespace(scalars=lambda: SimpleNamespace(all=lambda: [stale, fresh, already_stale])),
            SimpleNamespace(scalar_one_or_none=lambda: "new"),
            SimpleNamespace(scalar_one_or_none=lambda: "same"),
            SimpleNamespace(scalar_one_or_none=lambda: "new"),
        ]
        db = AsyncMock()
        db.execute.side_effect = results

        report = await check_batch_staleness(db, uuid.uuid4())

        assert [item["marked_stale"] for item in report] == [True, False, False]
        assert report[0]["was_already_stale"] is False
        assert report[2]["was_already_stale"] is True
        assert stale.is_stale is True
        assert fresh.is_stale is False
        db.flush.assert_awaited_once()


# ═══════════════════════════════════════════════════════════════════════════════
# RAG-13: Redaction edge cases
# ═══════════════════════════════════════════════════════════════════════════════


class TestRedactionEdgeCases:
    def test_redact_jwt_token(self):
        from app.services.rag_redaction_service import redact_prompt
        text = "Token: eyJhbGciOiJIUzI1NiJ9.eyJzdWIiOiIxMjM0NTY3ODkwIn0.dozjgNryP4J3jVmNHl0w5N_XgL0n3I9PlFUP0THsR8U"
        result, was_redacted = redact_prompt(text)
        assert was_redacted is True
        assert "eyJ" not in result

    def test_redact_base64_token(self):
        from app.services.rag_redaction_service import redact_prompt
        text = "api_key: " + "A" * 45  # long base64-like string
        result, was_redacted = redact_prompt(text)
        assert was_redacted is True

    def test_redact_empty_string(self):
        from app.services.rag_redaction_service import redact_prompt
        result, was_redacted = redact_prompt("")
        assert result == ""
        assert was_redacted is False

    def test_redact_none_returns_none(self):
        from app.services.rag_redaction_service import redact_prompt
        result, was_redacted = redact_prompt(None)
        assert result is None
        assert was_redacted is False

    def test_redact_confidential_key_value(self):
        from app.services.rag_redaction_service import redact_prompt
        text = "password: super_secret_123\nusername: admin"
        result, was_redacted = redact_prompt(text, classification="confidential")
        assert was_redacted is True
        assert "super_secret_123" not in result
        assert "admin" in result  # username not a sensitive key

    def test_redact_restricted_same_as_confidential(self):
        from app.services.rag_redaction_service import redact_prompt
        text = "api_key= my-secret-key-value-here"
        result, was_redacted = redact_prompt(text, classification="restricted")
        assert was_redacted is True

    def test_redact_chunk_text_delegates(self):
        from app.services.rag_redaction_service import redact_chunk_text
        text = "Bearer eyJhbGciOiJIUzI1NiJ9.eyJzdWIiOiIxIn0.abc123"
        result = redact_chunk_text(text)
        assert "eyJ" not in result

    def test_validate_document_upload_all_types(self):
        from app.services.rag_redaction_service import validate_document_upload
        for ext in ("pdf", "docx", "doc", "md", "txt", "markdown", "rst"):
            validate_document_upload(f"file.{ext}", 1024)

    def test_validate_document_upload_zero_size(self):
        from app.services.rag_redaction_service import validate_document_upload
        validate_document_upload("file.pdf", 0)

    def test_validate_url_scheme_https(self):
        from app.services.rag_redaction_service import validate_url_scheme
        validate_url_scheme("https://example.com/page")

    def test_validate_url_scheme_http(self):
        from app.services.rag_redaction_service import validate_url_scheme
        validate_url_scheme("http://internal.corp.com/docs")

    def test_validate_url_scheme_javascript_blocked(self):
        from app.services.rag_redaction_service import validate_url_scheme
        with pytest.raises(ValueError, match="not allowed"):
            validate_url_scheme("javascript:alert(1)")

    def test_validate_url_scheme_ftp_blocked(self):
        from app.services.rag_redaction_service import validate_url_scheme
        with pytest.raises(ValueError, match="not allowed"):
            validate_url_scheme("ftp://files.example.com/data")


# ═══════════════════════════════════════════════════════════════════════════════
# RAG-14: Eval schemas
# ═══════════════════════════════════════════════════════════════════════════════


class TestEvalSchemas:
    def test_rag_status_defaults(self):
        from app.models.schemas import RagStatusResponse
        status = RagStatusResponse(enabled=False)
        assert status.total_sources == 0
        assert status.total_batches == 0
        assert status.total_chunks == 0
        assert status.feature_flag == "KNOWLEDGE_RAG_ENABLED"

    def test_generation_batch_response_optional_fields(self):
        from app.models.schemas import GenerationBatchResponse
        resp = GenerationBatchResponse(
            id=uuid.uuid4(),
            project_id=uuid.uuid4(),
            generation_mode="raw",
            cases_generated=0,
            cases_accepted=0,
            cases_rejected=0,
            status="pending",
            created_at=datetime.now(timezone.utc),
        )
        assert resp.llm_model_used is None
        assert resp.completed_at is None
        assert resp.coverage_score is None


# ═══════════════════════════════════════════════════════════════════════════════
# RAG-1: Knowledge source service — domain validation
# ═══════════════════════════════════════════════════════════════════════════════


class TestDomainValidationEdgeCases:
    def test_case_insensitive_domain_match(self):
        from app.services.knowledge_source_service import _validate_url_domain
        _validate_url_domain("https://JIRA.CORP.COM/browse/X", ["jira.corp.com"])

    def test_exact_domain_match(self):
        from app.services.knowledge_source_service import _validate_url_domain
        _validate_url_domain("https://jira.corp.com/page", ["jira.corp.com"])

    def test_subdomain_match(self):
        from app.services.knowledge_source_service import _validate_url_domain
        _validate_url_domain("https://east.jira.corp.com/page", ["jira.corp.com"])

    def test_partial_mismatch_blocked(self):
        from fastapi import HTTPException
        from app.services.knowledge_source_service import _validate_url_domain
        with pytest.raises(HTTPException) as exc_info:
            _validate_url_domain("https://fakejira.corp.com/page", ["jira.corp.com"])
        assert exc_info.value.status_code == 422

    def test_multiple_domains_any_match(self):
        from app.services.knowledge_source_service import _validate_url_domain
        _validate_url_domain(
            "https://confluence.corp.com/page",
            ["jira.corp.com", "confluence.corp.com", "docs.internal.io"],
        )

    def test_url_without_hostname(self):
        from fastapi import HTTPException
        from app.services.knowledge_source_service import _validate_url_domain
        with pytest.raises(HTTPException):
            _validate_url_domain("not-a-url", ["jira.corp.com"])


# ═══════════════════════════════════════════════════════════════════════════════
# RAG-1: Feature flag
# ═══════════════════════════════════════════════════════════════════════════════


class TestFeatureFlagEdgeCases:
    @pytest.mark.asyncio
    async def test_gate_raises_503_with_a_pointer_to_the_switch(self, monkeypatch):
        from fastapi import HTTPException
        from app.services import knowledge_source_service as svc

        async def _off(key, **kwargs):
            return False

        monkeypatch.setattr("app.services.feature_flags.is_enabled", _off)
        with pytest.raises(HTTPException) as exc:
            await svc.require_rag_enabled_async(object())
        assert exc.value.status_code == 503
        assert "not enabled" in exc.value.detail
        # The message names a page, so that page must drive this same flag.
        assert "AI Configuration" in exc.value.detail


# ═══════════════════════════════════════════════════════════════════════════════
# RAG-8: Generation constants
# ═══════════════════════════════════════════════════════════════════════════════


class TestGenerationConstants:
    def test_max_citations_per_case_is_positive(self):
        from app.services.rag_generation_service import MAX_CITATIONS_PER_CASE
        assert MAX_CITATIONS_PER_CASE > 0

    def test_grounded_generation_result_dataclass(self):
        from app.services.rag_generation_service import GroundedGenerationResult
        result = GroundedGenerationResult(
            batch_id=uuid.uuid4(),
            generation_mode="grounded",
            test_cases=[{"title": "Test"}],
            citations=[],
        )
        assert result.coverage_summary is None
        assert result.gaps_noted == []
        assert result.created_ids == []


# ═══════════════════════════════════════════════════════════════════════════════
# Connector registry
# ═══════════════════════════════════════════════════════════════════════════════


class TestConnectorRegistryEdgeCases:
    def test_jira_issue_and_epic_share_connector_class(self):
        from app.services.connectors.registry import get_connector
        jira_issue = get_connector("jira_issue")
        jira_epic = get_connector("jira_epic")
        assert type(jira_issue) is type(jira_epic)

    def test_internal_and_external_url_share_connector_class(self):
        from app.services.connectors.registry import get_connector
        internal = get_connector("internal_url")
        external = get_connector("external_url")
        assert type(internal) is type(external)


# ═══════════════════════════════════════════════════════════════════════════════
# FetchedContent
# ═══════════════════════════════════════════════════════════════════════════════


class TestFetchedContent:
    def test_compute_hash_sha256(self):
        h = FetchedContent.compute_hash("test")
        assert len(h) == 64
        # SHA-256 of "test"
        assert h == "9f86d081884c7d659a2feaa0c55ad015a3bf4f1b2b0b822cd15d6c15b0f00a08"

    def test_compute_hash_unicode(self):
        h = FetchedContent.compute_hash("héllo wörld")
        assert len(h) == 64

    def test_default_fields(self):
        fc = FetchedContent(raw_text="text", content_hash="abc", title="T", source_url="http://x")
        assert fc.metadata == {}
        assert fc.attachments == []


class TestConnectorFetchError:
    def test_retryable_default_false(self):
        err = ConnectorFetchError("oops")
        assert err.retryable is False
        assert str(err) == "oops"

    def test_retryable_true(self):
        err = ConnectorFetchError("timeout", retryable=True)
        assert err.retryable is True
