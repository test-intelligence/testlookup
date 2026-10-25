"""W13 TCL-04: import test cases from CSV / Excel (``POST /cases/import``).

The real parser and service over a recording session:
1. a dry run of 20 rows (3 bad) -> 17 ``create``, 3 ``invalid`` with sentences and
   row numbers; the session saw NO add / flush / commit;
2. ``Status=approved`` -> ``create`` + its warning; applied, the case is a Draft
   and an ``imported`` audit row names the override;
3. ``=SUM(A1:A2)`` is stored as typed, an export's ``'=SUM(A1:A2)`` loses ONE
   ``'`` (with its warning), and the export writes ``'=SUM(A1:A2)`` again as text;
4. duplicates of an existing case and of an earlier row; ``p0``; severity
   ``medium`` kept with a warning; ``Automated=maybe`` invalid;
5. the caps: 2,001 rows; 5 MB + 1 while reading; an .xlsx expanding past 50 MB
   before ``load_workbook``; a ``.txt``; no Title column;
6. one row's create failing -> that row invalid, the others created, ONE
   ``test_case.imported`` row; zero created -> no row and no commit;
7. TESTER allowed, VIEWER 403, a non-member 403 before the file is read.
"""
from __future__ import annotations

import io
import uuid
import zipfile
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from fastapi import FastAPI, HTTPException
from httpx import ASGITransport, AsyncClient

from app.core import deps
from app.core import test_library_vocab as v
from app.core.deps import get_current_active_user, get_db
from app.models.postgres import ManagedTestCase, Project, ProjectActivityEvent, UserRole
from app.models.postgres import TestCaseAuditLog as AuditRow
from app.routers import test_management_library as routes
from app.services import test_case_import_service as svc
from app.services import test_management_service

PROJECT = uuid.uuid4()
HEADER = "Title,Type,Priority,Severity,Objective,Steps,Expected Result,Automated,Tags,Status"


class _Result:
    def __init__(self, rows=()):
        self._rows = list(rows)

    def all(self):
        return list(self._rows)


class _Savepoint:
    def __init__(self, session):
        self.session = session

    async def __aenter__(self):
        self.mark = len(self.session.added)
        return self

    async def __aexit__(self, exc_type, exc, tb):
        if exc_type is not None:
            del self.session.added[self.mark:]          # the savepoint's rows roll back
        return False


class _Session:
    """Records every add / flush / commit; answers the duplicate read with ``existing``."""

    def __init__(self, existing=()):
        self.existing = list(existing)
        self.added: list = []
        self.flushes = 0
        self.commits = 0
        self.info: dict = {}
        self.execute = AsyncMock(side_effect=lambda *a, **k: _Result(self.existing))

    def add(self, obj):
        if isinstance(obj, ManagedTestCase) and obj.id is None:
            obj.id = uuid.uuid4()
        self.added.append(obj)

    async def flush(self):
        self.flushes += 1

    def begin_nested(self):
        return _Savepoint(self)

    async def commit(self):
        self.commits += 1

    async def get(self, model, ident):
        return SimpleNamespace(id=ident) if model is Project else None


def _actor(role=UserRole.TESTER):
    return SimpleNamespace(id=uuid.uuid4(), role=role.value, is_active=True, username="tess", full_name="Tess T",
                           email="t@x.test", is_service_account=False)


@pytest.fixture(autouse=True)
def member(monkeypatch):
    monkeypatch.setattr(deps, "get_accessible_project_ids", AsyncMock(return_value={PROJECT}))


def _csv(*lines: str) -> bytes:
    return ("\n".join([HEADER, *lines]) + "\n").encode("utf-8")


async def _run(content: bytes, *, dry_run=True, existing=(), name="cases.csv", session=None):
    db = session or _Session(existing)
    body = await svc.import_cases(db, project_id=PROJECT, actor=_actor(), file_name=name, content=content,
                                  dry_run=dry_run)
    return db, body


def _good(n: int) -> str:
    return f"Case number {n},e2e,high,major,Checks {n},Step 1: open; Step 2: pay {n},Paid,Yes,smoke; pay,"


async def test_a_dry_run_reports_every_row_and_writes_nothing():
    lines = [_good(n) for n in range(17)]
    lines.insert(3, "No,e2e,high,major,,,,No,,")                          # row 5: title too short
    lines.insert(8, "Bad priority,e2e,urgent,major,,,,No,,")               # row 10
    lines.insert(12, "Bad type,e2e!,high,major,,,,No,,")                    # row 14
    db, body = await _run(_csv(*lines))

    assert body["dry_run"] is True and body["format"] == "csv"
    assert body["summary"] == {"rows": 20, "create": 17, "duplicate": 0, "invalid": 3, "created": 0}
    invalid = [(r["row"], r["reasons"]) for r in body["rows"] if r["outcome"] == "invalid"]
    assert invalid == [(5, [v.IMPORT_ROW_TITLE]), (10, [v.IMPORT_ROW_PRIORITY]), (14, [v.IMPORT_ROW_TYPE])]
    assert body["headers"]["recognised"][:3] == ["Title", "Type", "Priority"]
    assert body["headers"]["ignored"] == ["Status"] and body["headers"]["unknown"] == []
    first = body["rows"][0]
    assert first["outcome"] == "create" and first["title"] == "Case number 0"
    assert (db.added, db.flushes, db.commits) == ([], 0, 0)                  # a dry run writes nothing
    assert db.execute.await_count == 1                                       # the ONE duplicate read


async def test_a_status_is_ignored_with_a_warning_and_audited_on_apply():
    content = _csv("Approved before,e2e,high,major,,,,No,,approved")
    _db, dry = await _run(content)
    assert dry["rows"][0]["outcome"] == "create"
    assert dry["rows"][0]["warnings"] == [v.IMPORT_WARN_STATUS.format(value="approved")]

    db, applied = await _run(content, dry_run=False)
    assert applied["summary"]["created"] == 1 and applied["rows"][0]["outcome"] == "created"
    case = next(o for o in db.added if isinstance(o, ManagedTestCase))
    assert case.status == "draft" and case.version == 1 and case.dup_fingerprint
    assert applied["rows"][0]["case_id"] == case.id
    audits = [o for o in db.added if isinstance(o, AuditRow)]
    assert [a.action for a in audits] == ["created", "imported"]
    assert "Status 'approved' was not applied" in audits[1].details and "(row 2)" in audits[1].details
    assert db.commits == 0                                                   # the router commits


async def test_formula_text_is_data_and_the_export_quote_comes_off_once():
    from openpyxl import Workbook

    from app.routers.test_management_exports import _excel_text

    content = _csv("=SUM(A1:A2),e2e,high,major,,,,No,,", "'=SUM(A1:A2) again,e2e,high,major,,,,No,,",
                   "''=x,e2e,high,major,,,,No,,")
    _db, body = await _run(content)
    titles = [(r["title"], r["warnings"]) for r in body["rows"]]
    assert titles[0] == ("=SUM(A1:A2)", [])
    assert titles[1] == ("=SUM(A1:A2) again", [v.IMPORT_WARN_QUOTE.format(column="Title")])
    assert titles[2] == ("''=x", [])                                         # only an export's ONE quote
    exported = _excel_text("=SUM(A1:A2)")
    cell = Workbook().active.cell(row=1, column=1, value=exported)
    assert exported == "'=SUM(A1:A2)" and cell.data_type == "s"


async def test_duplicates_aliases_kept_severity_and_a_bad_automated_flag():
    existing_fp = svc._fingerprint("Already here", None, None, None)
    existing = [SimpleNamespace(id=uuid.uuid4(), title="Already here", objective=None, steps=None,
                                expected_result=None, dup_fingerprint=existing_fp)]
    content = _csv("Already here,e2e,high,major,,,,No,,", "Twice in the file,e2e,p0,medium,,,,No,,",
                   "Twice in the file,e2e,high,major,,,,No,,", "Maybe automated,e2e,high,major,,,,maybe,,")
    _db, body = await _run(content, existing=existing)
    rows = body["rows"]
    assert rows[0]["outcome"] == "duplicate" and rows[0]["reasons"] == [v.IMPORT_DUPLICATE_EXISTING]
    assert rows[0]["duplicate_of"] == {"case_id": existing[0].id, "title": "Already here"}
    assert rows[1]["outcome"] == "create"
    assert rows[1]["warnings"] == [v.IMPORT_WARN_SEVERITY.format(value="medium")]
    assert rows[2]["outcome"] == "duplicate" and rows[2]["duplicate_of"] == {"row": 3}
    assert rows[2]["reasons"] == [v.IMPORT_DUPLICATE_ROW.format(row=3)]
    assert rows[3]["outcome"] == "invalid" and rows[3]["reasons"] == [v.IMPORT_ROW_AUTOMATED]
    check = svc.check_row(2, {"title": "Alias", "priority": "P0"})
    assert check.fields["priority"] == "critical" and not check.reasons


def test_steps_read_the_exports_format_or_one_per_line():
    assert [s["action"] for s in svc.parse_steps("Step 1: open; Step 2: pay; with a semicolon")] == [
        "open", "pay; with a semicolon"]
    assert [s["action"] for s in svc.parse_steps("open\n\n pay \n")] == ["open", "pay"]
    assert svc.parse_steps("Step 1: open")[0] == {"step_number": 1, "action": "open", "expected_result": ""}


# ── 5. the caps, before anything is parsed further ───────────────────────────


def _refused(call, status_code, detail):
    with pytest.raises(HTTPException) as exc:
        call()
    assert (exc.value.status_code, exc.value.detail) == (status_code, detail)


def test_more_than_2000_rows_is_refused_and_2000_is_not():
    rows = [f"Case {n},e2e,high,major,,,,No,," for n in range(2001)]
    _refused(lambda: svc.parse_upload("big.csv", _csv(*rows)), 422, v.IMPORT_TOO_MANY_ROWS)
    assert len(svc.parse_upload("ok.csv", _csv(*rows[:2000])).rows) == 2000
    assert len(svc.parse_upload("blank.csv", _csv(*rows[:3], ",,,,,,,,,", "", *rows[3:5])).rows) == 5


class _Upload:
    def __init__(self, total: int, chunk: int = 64 * 1024):
        self.left, self.chunk, self.reads = total, chunk, 0

    async def read(self, n: int = -1) -> bytes:
        self.reads += 1
        size = min(self.left, n if n > 0 else self.chunk)
        self.left -= size
        return b"x" * size


async def test_an_upload_over_5_mb_is_refused_while_reading():
    upload = _Upload(v.IMPORT_MAX_BYTES + 1 + 10 * 1024 * 1024)            # 15 MB on offer
    with pytest.raises(HTTPException) as exc:
        await svc.read_bounded(upload)
    assert (exc.value.status_code, exc.value.detail) == (413, v.IMPORT_TOO_LARGE)
    assert upload.reads == v.IMPORT_MAX_BYTES // (64 * 1024) + 1            # stopped at the cap
    assert len(await svc.read_bounded(_Upload(v.IMPORT_MAX_BYTES))) == v.IMPORT_MAX_BYTES


def test_a_workbook_expanding_past_50_mb_is_refused_before_load_workbook(monkeypatch):
    import openpyxl

    def _never(*args, **kwargs):
        raise AssertionError("load_workbook must not be reached")

    monkeypatch.setattr(openpyxl, "load_workbook", _never)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/sheet1.xml", b"\0" * (v.IMPORT_MAX_XLSX_UNCOMPRESSED + 1))
    content = buffer.getvalue()
    assert len(content) < v.IMPORT_MAX_BYTES                                # small on the wire
    _refused(lambda: svc.parse_upload("bomb.xlsx", content), 422, v.IMPORT_XLSX_TOO_BIG)


def test_the_format_and_header_sentences():
    _refused(lambda: svc.parse_upload("cases.txt", b"Title\nx"), 400, v.IMPORT_NOT_A_FILE)
    _refused(lambda: svc.parse_upload("fake.xlsx", b"Title\nnot a zip"), 400, v.IMPORT_NOT_A_FILE)
    _refused(lambda: svc.parse_upload("cases.csv", b"Name,Steps\nx,y\n"), 422, v.IMPORT_NO_TITLE_COLUMN)
    _refused(lambda: svc.parse_upload("cases.csv", b"\xff\xfe\x00bad"), 422, v.IMPORT_UNREADABLE)
    _refused(lambda: svc.parse_upload("empty.csv", b""), 422, v.IMPORT_NO_TITLE_COLUMN)


def test_an_xlsx_reads_numbers_dates_and_the_first_sheet():
    from datetime import date

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Title", "Priority", "Tags", "Owner"])
    sheet.append(["Excel case", "P1", 42, "nobody"])
    sheet.append(["Dated", "low", date(2026, 10, 24), None])
    workbook.create_sheet("Other").append(["Title"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    parsed = svc.parse_upload("cases.XLSX", buffer.getvalue())
    assert parsed.format == "xlsx" and parsed.unknown == ["Owner"]
    assert parsed.rows == [(2, {"title": "Excel case", "priority": "P1", "tags": "42"}),
                           (3, {"title": "Dated", "priority": "low", "tags": "2026-10-24"})]


# ── 6. a failing row, and the activity row ───────────────────────────────────


async def test_one_failing_create_leaves_the_others_and_one_activity_row(monkeypatch):
    real = test_management_service.create_managed_test_case

    async def _create(db, payload, actor):
        if payload.title == "Case number 1":
            raise HTTPException(status_code=404, detail="Project gone")
        return await real(db, payload, actor)

    monkeypatch.setattr(test_management_service, "create_managed_test_case", _create)
    db, body = await _run(_csv(_good(0), _good(1), _good(2)), dry_run=False)
    assert [r["outcome"] for r in body["rows"]] == ["created", "invalid", "created"]
    assert body["rows"][1]["reasons"] == [v.IMPORT_SAVE_FAILED.format(detail="Project gone")]
    assert body["summary"] == {"rows": 3, "create": 3, "duplicate": 0, "invalid": 1, "created": 2}
    events = [o for o in db.added if isinstance(o, ProjectActivityEvent)]
    assert [(e.event_type, e.entity_type, e.entity_label) for e in events] == [
        ("test_case.imported", "test_library", "cases.csv")]
    assert events[0].context == {"count": 2, "duplicate": 0, "invalid": 1}


async def test_nothing_created_stages_no_activity_row():
    db, body = await _run(_csv("No,e2e,high,major,,,,No,,"), dry_run=False)
    assert body["summary"]["created"] == 0
    assert not [o for o in db.added if isinstance(o, ProjectActivityEvent)]


# ── 7. the route: roles, scope, the commit ───────────────────────────────────


def _client(session, user):
    app = FastAPI()
    app.include_router(routes.router, prefix="/api/v1/test-management")

    async def _db():
        yield session

    app.dependency_overrides[get_db] = _db
    app.dependency_overrides[get_current_active_user] = lambda: user
    return AsyncClient(transport=ASGITransport(app=app), base_url="http://t")


async def _post(session, user, *, dry_run, content=None, project=PROJECT):
    files = {"file": ("cases.csv", content or _csv(_good(0)), "text/csv")}
    async with _client(session, user) as client:
        return await client.post("/api/v1/test-management/cases/import",
                                 params={"project_id": str(project), "dry_run": str(dry_run).lower()}, files=files)


async def test_a_tester_previews_then_applies_with_one_commit():
    session = _Session()
    preview = await _post(session, _actor(), dry_run=True)
    assert preview.status_code == 200, preview.text
    assert preview.json()["summary"]["create"] == 1 and session.commits == 0
    applied = await _post(session, _actor(), dry_run=False)
    assert applied.status_code == 200 and applied.json()["summary"]["created"] == 1
    assert session.commits == 1


async def test_nothing_created_is_not_committed():
    session = _Session()
    res = await _post(session, _actor(), dry_run=False, content=_csv("No,e2e,high,major,,,,No,,"))
    assert res.status_code == 200 and session.commits == 0


async def test_a_viewer_and_a_non_member_are_refused_before_the_file_is_read(monkeypatch):
    read = AsyncMock(side_effect=AssertionError("the file must not be read"))
    monkeypatch.setattr(routes, "read_bounded", read)
    viewer = await _post(_Session(), _actor(UserRole.VIEWER), dry_run=True)
    assert (viewer.status_code, viewer.json()["detail"]) == (403, "Requires at least TESTER role")
    outsider = await _post(_Session(), _actor(), dry_run=True, project=uuid.uuid4())
    assert (outsider.status_code, outsider.json()["detail"]) == (403, "You do not have access to this project")
    read.assert_not_awaited()


async def test_the_dry_run_is_the_default():
    session = _Session()
    files = {"file": ("cases.csv", _csv(_good(0)), "text/csv")}
    async with _client(session, _actor()) as client:
        res = await client.post("/api/v1/test-management/cases/import", params={"project_id": str(PROJECT)},
                                files=files)
    assert res.json()["dry_run"] is True and session.commits == 0 and not session.added
